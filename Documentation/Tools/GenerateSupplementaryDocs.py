#!/usr/bin/env python3
"""Render the CMake and FemboyScript manuals to DOCX and PDF.

The Markdown files are the source of truth. This script intentionally supports a
small Markdown subset: headings, paragraphs, bullets, numbered lists, fenced
code and pipe tables.
"""

from __future__ import annotations

import argparse
import datetime as dt
import os
import pathlib
import re
import subprocess
import sys

sys.dont_write_bytecode = True

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import DocumentConvert as docconvert
from DocVersion import VERSION as DOC_VERSION

ROOT = pathlib.Path(__file__).resolve().parents[2]
DOCS = ROOT / "Documentation"
TODAY = dt.date.today().isoformat()

DOCUMENTS = [
    {
        "stem": "CMakeProjectManagement",
        "title_en": "CMake Project and Engine Management",
        "title_ru": "Управление проектом и движком через CMake",
    },
    {
        "stem": "FemboyScript",
        "title_en": "FemboyScript",
        "title_ru": "FemboyScript",
    },
]


def set_cell_background(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def add_inline_runs(paragraph, text: str) -> None:
    text = text.replace("**", "")
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.Size = Pt(9)
            run.font.color.rgb = RGBColor(0x8B, 0x1A, 0x4A)
        else:
            paragraph.add_run(part)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


def render_markdown(md_path: pathlib.Path, docx_path: pathlib.Path, title: str, language: str) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = document.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.Size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)

    for level in range(1, 4):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Segoe UI"
        style.font.color.rgb = RGBColor(0x8B, 0x1A, 0x4A)

    # Title block.
    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.name = "Segoe UI"
    title_run.font.Size = Pt(26)
    title_run.font.color.rgb = RGBColor(0x8B, 0x1A, 0x4A)
    subtitle = document.add_paragraph()
    subtitle.add_run(f"FemboyEngine2D {DOC_VERSION}  ·  {language}  ·  {TODAY}").font.Size = Pt(9)
    document.add_paragraph()

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            index += 1
            code_paragraph = document.add_paragraph()
            code_paragraph.style = document.styles["No Spacing"]
            code_run = code_paragraph.add_run("\n".join(code_lines))
            code_run.font.name = "Consolas"
            code_run.font.Size = Pt(8)
            code_paragraph.paragraph_format.left_indent = Cm(0.5)
            code_paragraph.paragraph_format.space_after = Pt(6)
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?[-: ]+\|", lines[index + 1]):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", cell or "-") for cell in row):
                    rows.append(row)
                index += 1
            if rows:
                table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    for column_index, value in enumerate(row):
                        cell = table.cell(row_index, column_index)
                        cell.text = ""
                        paragraph = cell.paragraphs[0]
                        add_inline_runs(paragraph, value)
                        if row_index == 0:
                            for run in paragraph.runs:
                                run.bold = True
                            set_cell_background(cell, "F4E3EA")
                document.add_paragraph()
            continue

        if stripped.startswith("#"):
            level = min(len(stripped) - len(stripped.lstrip("#")), 3)
            document.add_heading(stripped[level:].strip(), level=level)
            index += 1
            continue

        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, stripped[2:])
            index += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_inline_runs(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            index += 1
            continue

        if stripped:
            paragraph = document.add_paragraph()
            add_inline_runs(paragraph, stripped)
        index += 1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"FemboyEngine2D {DOC_VERSION}  ·  {title}  ·  ").font.Size = Pt(8)
    add_page_number(footer)

    document.core_properties.title = title
    document.core_properties.author = "FemboyEngine2D contributors"
    document.core_properties.subject = f"FemboyEngine2D {DOC_VERSION} documentation"
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(docx_path)


def export_pdf(docx_path: pathlib.Path, pdf_path: pathlib.Path, node: str, lo_cli: str) -> bool:
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    if pdf_path.exists():
        pdf_path.unlink()
    if not pathlib.Path(lo_cli).is_file():
        print(f"pdf: LibreOffice Kit CLI not found: {lo_cli}", file=sys.stderr)
        return False
    completed = subprocess.run(
        [node, lo_cli, "convert", "--input", str(docx_path), "--output", str(pdf_path)],
        check=False,
        capture_output=True,
        text=True,
    )
    if completed.returncode != 0:
        print(completed.stdout[-2000:], file=sys.stderr)
        print(completed.stderr[-2000:], file=sys.stderr)
        return False
    return pdf_path.is_file()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--node", default=os.environ.get("FBY_NODE", "node"))
    parser.add_argument(
        "--lo-cli",
        default=os.environ.get("FBY_LO_CLI", ""),
    )
    parser.add_argument("--skip-pdf", action="store_true")
    args = parser.parse_args()

    failures = 0
    for document in DOCUMENTS:
        for language, title_key in (("en", "title_en"), ("ru", "title_ru")):
            md_path = DOCS / "md" / language / f"{document['stem']}_{language.upper()}.md"
            docx_path = DOCS / "docx" / language / f"{document['stem']}_{language.upper()}.docx"
            pdf_path = DOCS / "pdf" / language / f"{document['stem']}_{language.upper()}.pdf"
            if not md_path.is_file():
                print(f"missing Markdown: {md_path}", file=sys.stderr)
                failures += 1
                continue
            render_markdown(md_path, docx_path, document[title_key], language.upper())
            print(f"docx: {docx_path.relative_to(ROOT)}")
            if not args.skip_pdf and not export_pdf(docx_path, pdf_path, args.node, args.lo_cli):
                failures += 1
            elif not args.skip_pdf:
                print(f"pdf: {pdf_path.relative_to(ROOT)}")
                djvu_path = DOCS / "djvu" / language / f"{document['stem']}_{language.upper()}.djvu"
                docconvert.BuildDjvu(pdf_path, djvu_path)
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
